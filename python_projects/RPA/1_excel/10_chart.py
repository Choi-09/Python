from openpyxl import load_workbook

wb = load_workbook("rpa_test5.xlsx")
ws = wb.active

from openpyxl.chart import BarChart, LineChart, Reference # Reference도 같이 추가
# # Bar chart 추가
# # 1. 어떤 데이터를 차트로 만들지 value 설정
# bar_value = Reference(ws, min_row = 2, max_row = 11, min_col = 2, max_col = 3)  # B2 ~ C11까지의 데이터로 차트 생성

# # 2. 차트 종류 설정 (Bar, Line, Pie, ...)
# bar_chart = BarChart()

# # 3. 차트에 데이터 추가
# bar_chart.add_data(bar_value)

# # 4. 차트 넣을 위치 정하기
# ws.add_chart(bar_chart, "E1") 

# wb.save("rpa_test5_bar_chart.xlsx")

# Line chart 추가 
# 1. 데이터 value 설정
line_value = Reference(ws, min_row = 1, max_row = 11, min_col = 2, max_col = 3) # min_row = 1: 컬럼명을 포함해서 차트 생성

# 2. 차트 종류 설정 
line_chart = LineChart()

# 3. 차트에 데이터 추가
line_chart.add_data(line_value, titles_from_data = True) # 레이블 이름을 컬럼명으로 지정

# 4. 차트 타이틀 지정
line_chart.title = "성적표"

# 5. 미리 정의된 스타일을 적용, 사용자가 개별 지정도 가능
line_chart.style = 10

# 6. 차트 Y축 제목 설정
line_chart.y_axis.title = "점수"

# 7. 차트 X축 제목 설정
line_chart.x_axis.title = "번호"

# 8. 차트 넣을 위치 설정
ws.add_chart(line_chart, "E1")

wb.save("rpa_test5_line_chart.xlsx")